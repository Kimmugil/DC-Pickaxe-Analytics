"""기존 엑셀(구글시트 내려받기) → Supabase 이관. 기본은 점검만(dry-run), --apply 를 붙여야 실제 저장."""
import glob, os, re, sys, collections
from datetime import datetime, timezone, timedelta
from urllib.parse import urlparse, parse_qs
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")

SRC = sys.argv[1]
APPLY = "--apply" in sys.argv
KST = timezone(timedelta(hours=9))
TYPE_BY_PATH = {"mgallery/board": "minor", "board": "major", "mini/board": "mini"}
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def clean(x):
    return CTRL.sub("", str(x)) if x is not None else None

def num(x):
    s = str(x or "0").split("/")[0].replace(",", "").strip()
    return int(s) if s.lstrip("-").isdigit() else 0

def parse_dt(s):
    s = str(s).strip()
    for fmt, known in (("%Y-%m-%d %H:%M", True), ("%Y-%m-%d", False)):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=KST), known
        except ValueError:
            pass
    return None, False

NAMES = {"dusdns": "연운(개봉)", "wherewindsmeets": "연운", "dfip": "던파IP", "azurpromilia": "아주르 프로밀리아", "projectnike": "승리의 여신 니케"}
master = {}
ws = load_workbook(os.path.join(SRC, "00_DC-Pickaxe_Master.xlsx"), read_only=True, data_only=True).worksheets[0]
for r in list(ws.iter_rows(values_only=True))[1:]:
    if r and r[1]: master[r[1]] = r[0]

galleries, posts, problems = {}, {}, collections.Counter()
for f in sorted(glob.glob(os.path.join(SRC, "[0-9][0-9]_*.xlsx"))):
    if os.path.basename(f).startswith("00_"): continue
    ws = load_workbook(f, read_only=True, data_only=True).worksheets[0]
    gid = None; n_file = 0
    for r in ws.iter_rows(values_only=True):
        if not r or r[0] is None or str(r[0]).strip() in ("", "글번호"): continue
        if not str(r[0]).strip().isdigit(): problems["글번호 비정상"] += 1; continue
        link = str(r[5] or "")
        q = parse_qs(urlparse(link).query)
        g = (q.get("id") or [None])[0]
        if not g: problems["링크에서 갤러리ID 못찾음"] += 1; continue
        if gid is None:
            gid = g
            path = urlparse(link).path.replace("/view/", "").strip("/")
            name = master.get(gid) or NAMES.get(gid) or gid
            galleries[gid] = (name, TYPE_BY_PATH.get(path, "minor"), gid in master)
        dt, known = parse_dt(r[4])
        if dt is None: problems["날짜 해석 실패"] += 1; continue
        key = (gid, int(str(r[0]).strip()))
        row = (gid, key[1], clean(r[1]) or "", clean(r[2]), clean(r[3]), dt, known, link,
                      num(r[6]), num(r[7]), num(r[8]))
        if key in posts:
            problems["중복 글번호(숫자 큰 건 유지)"] += 1
            if row[9] <= posts[key][9]: continue   # 조회수가 더 큰(나중에 갱신된) 쪽을 남김
        else:
            n_file += 1
        posts[key] = row
    print(f"{os.path.basename(f)[:40]:42} {gid or '?':14} {n_file:>7}건")

print("\n갤러리:", {k: v for k, v in galleries.items()})
print("총 게시글:", len(posts), "| 문제:", dict(problems))
print("날짜만 있는 글(시각 모름):", sum(1 for p in posts.values() if not p[6]))
if not APPLY:
    print("\n[dry-run] DB에 아무것도 저장하지 않았습니다."); sys.exit()

import psycopg
url = next(l.split("=", 1)[1].strip() for l in open(".env", encoding="utf-8") if l.startswith("DATABASE_URL="))
with psycopg.connect(url, connect_timeout=15) as c:
    with c.transaction():
        for gid, (name, typ, active) in galleries.items():
            c.execute("insert into galleries(gallery_id,name,gallery_type,is_active) values (%s,%s,%s,%s) "
                      "on conflict (gallery_id) do nothing", (gid, name, typ, active))
        with c.cursor() as cur, cur.copy(
            "copy posts (gallery_id,post_no,title,body,author,posted_at,posted_time_known,url,"
            "comment_count,view_count,recommend_count,metrics_updated_at,refresh_stage) from stdin") as cp:
            for p in posts.values():
                cp.write_row(p + (None, 2))   # 이관 글은 이미 갱신 완료 처리(stage 2)
    n = c.execute("select count(*) from posts").fetchone()[0]
    print("저장 완료. DB posts 행 수:", n)
