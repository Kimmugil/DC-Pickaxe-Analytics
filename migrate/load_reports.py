"""Analytics 엑셀 → reports / settings 이관. --apply 없으면 점검만."""
import ast, calendar, json, os, sys, collections
from datetime import date
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")
SRC, APPLY = sys.argv[1], "--apply" in sys.argv

def sheet(path, name):
    ws = load_workbook(path, read_only=True, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    return [{k: (r[i] if i < len(r) else None) for i, k in enumerate(h) if k} for r in rows[1:] if r and any(c is not None for c in r)]

def js(v):
    if v is None: return None
    if isinstance(v, (list, dict)): return v
    s = str(v).strip()
    if s in ("", "nan", "None"): return None
    for f in (json.loads, ast.literal_eval):
        try: return f(s)
        except Exception: pass
    return None

def num(v, cast=float):
    try: return cast(float(v))
    except Exception: return None

def txt(v):
    s = None if v is None else str(v).strip()
    return s or None

A = os.path.join(SRC, "Analytics", "DC-Pickaxe Analytics.xlsx")
reports = {}
def put(key, row): reports[key] = row   # 같은 키가 여러 번(재실행)이면 마지막 것 유지

stats = collections.Counter()
for r in sheet(A, "daily_issues"):
    d = date.fromisoformat(str(r["date"])[:10])
    put(("day", d, r["gallery_id"]), dict(
        period_type="day", period_start=d, period_end=d, gallery_id=r["gallery_id"],
        post_count=num(r["posts_total"], int), avg_7d=num(r["avg_7d"]), avg_same_weekday=num(r["avg_same_weekday"]),
        momentum_avg=num(r["momentum_avg"]), issue_score=num(r["issue_score"], int),
        has_issue=bool(num(r["has_issue"], int)), is_borderline=bool(num(r["is_borderline"], int)),
        headline=txt(r["headline"]), summary=txt(r["ai_summary"]), temperature_tag=txt(r["temperature_tag"]),
        top_cause=txt(r["issue_cause"]), category_scores=js(r["category_scores"]), major_issues=js(r["major_issues"]),
        sentiment={"positive": txt(r["sentiment_positive"]) or "", "negative": txt(r["sentiment_negative"]) or ""},
        top_posts=js(r["top_posts"]), keywords=js(r["keywords"]), daily_counts=None, run_id=txt(r["run_id"]), extra=None))
    stats["일간"] += 1
for r in sheet(A, "weekly_galleries"):
    s, e = date.fromisoformat(str(r["week_start"])[:10]), date.fromisoformat(str(r["week_end"])[:10])
    put(("week", s, r["gallery_id"]), dict(
        period_type="week", period_start=s, period_end=e, gallery_id=r["gallery_id"],
        post_count=num(r["total_posts"], int), avg_7d=None, avg_same_weekday=None, momentum_avg=None,
        issue_score=None, has_issue=False, is_borderline=False,
        headline=txt(r["headline"]), summary=txt(r["ai_summary"]), temperature_tag=txt(r["temperature_tag"]), top_cause=None,
        category_scores=js(r["category_scores"]), major_issues=js(r["major_issues"]),
        sentiment={"positive": txt(r["sentiment_positive"]) or "", "negative": txt(r["sentiment_negative"]) or ""},
        top_posts=js(r["top_posts"]), keywords=js(r["keywords"]), daily_counts=js(r["daily_counts"]), run_id=txt(r["run_id"]), extra=None))
    stats["주간"] += 1
for r in sheet(A, "weekly_overall"):
    s, e = date.fromisoformat(str(r["week_start"])[:10]), date.fromisoformat(str(r["week_end"])[:10])
    put(("week", s, None), dict(period_type="week", period_start=s, period_end=e, gallery_id=None, summary=txt(r["ai_summary"]), run_id=txt(r["run_id"])))
    stats["주간 종합"] += 1
for r in sheet(A, "monthly_issues"):
    y, m = map(int, str(r["month"]).split("-")); s = date(y, m, 1); e = date(y, m, calendar.monthrange(y, m)[1])
    put(("month", s, r["gallery_id"]), dict(
        period_type="month", period_start=s, period_end=e, gallery_id=r["gallery_id"],
        post_count=num(r["total_posts"], int), has_issue=False, is_borderline=False,
        headline=txt(r["headline"]), summary=txt(r["ai_summary"]), top_cause=txt(r["top_cause"]),
        category_scores=js(r["category_scores"]), major_issues=js(r["major_issues"]),
        sentiment={"positive": txt(r["sentiment_positive"]) or "", "negative": txt(r["sentiment_negative"]) or ""},
        top_posts=js(r["top_posts"]), keywords=js(r["keywords"]), daily_counts=js(r["daily_counts"]), run_id=txt(r["run_id"]),
        extra={"issue_days": num(r["issue_days"], int), "total_issue_score": num(r["total_issue_score"], int),
               "max_issue_score": num(r["max_issue_score"], int), "headlines": js(r["headlines"])}))
    stats["월간"] += 1
for r in sheet(A, "monthly_overall"):
    y, m = map(int, str(r["month"]).split("-")); s = date(y, m, 1)
    put(("month", s, None), dict(period_type="month", period_start=s, period_end=date(y, m, calendar.monthrange(y, m)[1]), gallery_id=None, summary=txt(r["ai_summary"]), run_id=txt(r["run_id"])))
    stats["월간 종합"] += 1

settings = {}
KEEP = {"app_title", "app_subtitle", "pm_name", "app_version", "announcement", "collection_note"}
for r in sheet(os.path.join(SRC, "00_DC-Pickaxe_Master.xlsx"), "config"):
    if r["key"] in KEEP and r["value"] is not None: settings[r["key"]] = str(r["value"])
for r in sheet(A, "ui_texts"):
    if r["key"] and r["value"] is not None: settings["ui." + str(r["key"])] = str(r["value"])   # 'secret' 시트(config)는 이관하지 않음

print("읽은 행:", dict(stats), "| 고유 리포트:", len(reports), "| 설정:", len(settings))
print("종류별:", dict(collections.Counter(k[0] for k in reports)))
if not APPLY: print("[dry-run] 저장하지 않음"); sys.exit()

import psycopg
from psycopg.types.json import Jsonb
cols = ["period_type","period_start","period_end","gallery_id","post_count","avg_7d","avg_same_weekday","momentum_avg","issue_score","has_issue","is_borderline","headline","summary","temperature_tag","top_cause","category_scores","major_issues","sentiment","top_posts","keywords","daily_counts","run_id","extra"]
url = next(l.split("=",1)[1].strip() for l in open(".env", encoding="utf-8") if l.startswith("DATABASE_URL="))
skipped = collections.Counter()
with psycopg.connect(url) as c, c.transaction():
    known = {r[0] for r in c.execute("select gallery_id from galleries")}
    for row in reports.values():
        if row.get("gallery_id") and row["gallery_id"] not in known: skipped[row["gallery_id"]] += 1; continue
        row = {"has_issue": False, "is_borderline": False, **row}
        vals = [row.get(k) for k in cols]
        vals = [Jsonb(v) if k in ("category_scores","major_issues","sentiment","top_posts","keywords","daily_counts","extra") and v is not None else v for k, v in zip(cols, vals)]
        c.execute(f"insert into reports ({','.join(cols)}) values ({','.join(['%s']*len(cols))}) "
                  "on conflict (period_type, period_start, coalesce(gallery_id,'')) do update set "
                  + ",".join(f"{k}=excluded.{k}" for k in cols[4:]) + ", updated_at=now()", vals)
    for k, v in settings.items():
        c.execute("insert into settings(key,value) values (%s,%s) on conflict (key) do update set value=excluded.value, updated_at=now()", (k, v))
    print("저장 완료 | 갤러리 미등록으로 건너뜀:", dict(skipped))
    for r in c.execute("select period_type, count(*) from reports group by 1 order by 1"): print(r)
    print("settings:", c.execute("select count(*) from settings").fetchone()[0])
